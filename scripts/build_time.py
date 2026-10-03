#!/usr/bin/env python3
"""Build screen-and-paper time series from NHK CSVs and MIC workbooks."""

from __future__ import annotations

import csv
import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "time" / "raw"
OUT = ROOT / "data" / "time" / "series.json"

SOUMU_YEAR = {
    "h24": 2012,
    "h25": 2013,
    "h26": 2014,
    "h27": 2015,
    "h28": 2016,
    "h29": 2017,
    "h30": 2018,
    "r1": 2019,
    "r2": 2020,
    "r3": 2021,
    "r4": 2022,
    "r5": 2023,
    "r6": 2024,
    "r7": 2025,
}


def hm_to_minutes(text: str) -> int:
    hour, minute = text.split(":")
    return int(hour) * 60 + int(minute)


def nhk_points(behavior: str, years: list[int]) -> list[dict]:
    points = []
    for year in years:
        path = RAW / f"nhk-{year}-4shihyo-all.csv"
        rows = list(csv.DictReader(path.read_text("utf-8-sig").splitlines()))
        hit = [
            row
            for row in rows
            if row["行動"] == behavior and row["曜日"] == "平日" and row["層"] == "国民全体"
        ]
        if not hit:
            continue
        points.append(
            {
                "year": year,
                "value": hm_to_minutes(hit[0]["全員平均時間量(時間:分)"]),
                "segment": "nhk",
                "source": f"nhk-{year}",
            }
        )
    return points


def sheet_rows(path: Path, year: int):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    name = book.sheetnames[0]
    if year >= 2024:
        for candidate in book.sheetnames:
            if "13-69" in candidate and "平日" in candidate and "休日" not in candidate:
                name = candidate
                break
    else:
        for candidate in book.sheetnames:
            if "平日" in candidate and "70" not in candidate and "13-69" not in candidate:
                name = candidate
                break
    sheet = book[name]
    rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
    book.close()
    avg_col = None
    for row in rows[:12]:
        for index, cell in enumerate(row):
            if cell and "平均" in str(cell) and "行為者" not in str(cell) and "標準" not in str(cell):
                avg_col = index
                break
        if avg_col is not None:
            break
    if avg_col is None:
        raise SystemExit(f"no average column in {path.name}")
    return name, rows, avg_col


def texts(row) -> list[str]:
    return [str(cell).replace("\n", "") if cell is not None else "" for cell in row]


def find_mean(rows, avg_col, predicate) -> float | None:
    for row in rows:
        if predicate(texts(row)):
            value = row[avg_col] if avg_col < len(row) else None
            if isinstance(value, (int, float)):
                return round(float(value), 1)
    return None


def soumu() -> dict[str, list[dict]]:
    series = {key: [] for key in ("tvLive", "tvRecorded", "internet", "newspaper", "radio", "magazine")}
    for code, year in SOUMU_YEAR.items():
        path = RAW / f"soumu-{code}-nikki.xlsx"
        sheet, rows, avg_col = sheet_rows(path, year)
        found = {
            "tvLive": find_mean(rows, avg_col, lambda t: any("テレビ（リアルタイム）" in cell for cell in t)),
            "tvRecorded": find_mean(rows, avg_col, lambda t: any(cell.startswith("録画したテレビ番組") for cell in t)),
            "internet": find_mean(rows, avg_col, lambda t: any("全てのインターネット利用" in cell for cell in t)),
            "newspaper": find_mean(rows, avg_col, lambda t: any("新聞閲読" in cell for cell in t)),
            "radio": find_mean(
                rows,
                avg_col,
                lambda t: any(cell.replace(" ", "").replace("　", "") in ("ラジオ", "ラジオ聴取") or "ラジオ聴取" in cell for cell in t),
            ),
            "magazine": find_mean(
                rows,
                avg_col,
                lambda t: any("雑誌（書籍" in cell for cell in t),
            ),
        }
        missing = [key for key, value in found.items() if value is None and key != "magazine"]
        if missing:
            raise SystemExit(f"{year} {sheet} missing {missing}")
        for key, value in found.items():
            if value is None:
                continue
            series[key].append(
                {
                    "year": year,
                    "value": value,
                    "segment": "age-13-69",
                    "source": f"soumu-{code}",
                    "sheet": sheet,
                }
            )
        print(year, sheet, "avg_col", avg_col, {k: found[k] for k in found})
    return series


def main():
    years = [1995, 2000, 2005, 2010, 2015, 2020, 2025]
    nhk = {
        "tv": ("テレビ", years),
        "radio": ("ラジオ", years),
        "newspaper": ("新聞", years),
        "magazineEarly": ("雑誌・マンガ", [1995, 2000]),
        "magazineLater": ("雑誌・マンガ・本", [2005, 2010, 2015, 2020, 2025]),
        "internet": ("趣味・娯楽・教養のインターネット", [2005, 2010, 2015]),
        "internetExVideo": ("趣味・娯楽・教養のインターネット（動画除く）", [2020, 2025]),
        "onlineVideo": ("インターネット動画", [2020, 2025]),
        "mediaContact": ("マスメディア接触", years),
    }
    built = {key: nhk_points(name, span) for key, (name, span) in nhk.items()}
    for key, points in built.items():
        if key == "internet" and len(points) != 3:
            raise SystemExit(f"nhk {key} {points}")
        if key in ("tv", "radio", "newspaper", "mediaContact") and len(points) != 7:
            raise SystemExit(f"nhk {key} incomplete {points}")
    mic = soumu()
    payload = {
        "question": "日本人は画面と紙にどれだけの時間を使ってきたか",
        "nhk": built,
        "soumu": mic,
    }
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print("wrote", OUT)


if __name__ == "__main__":
    main()
